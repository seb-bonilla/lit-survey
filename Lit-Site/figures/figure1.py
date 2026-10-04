"""Plot IZrO mobility vs carrier concentration with logarithmic resistivity colour."""
from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path

from openpyxl import load_workbook

Q = 1.602176634e-19  # elementary charge, C
ROOT = Path(__file__).resolve().parent.parent
CARRIER = 'Carrier concentration (cm^-3)'
MOBILITY = 'Mobility (cm^2/Vs)'


def positive_number(value):
    if isinstance(value, bool):
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) and value > 0 else None


def isoresistivity_mobility(carrier, resistivity_mohm_cm):
    """Single-carrier rho = 1/(q*n*mu); convert mΩ·cm to Ω·cm."""
    return 1000 / (Q * carrier * resistivity_mohm_cm)


def read_observations(workbook, derive_missing=False):
    values = load_workbook(workbook, read_only=True, data_only=True)
    formulas = load_workbook(workbook, read_only=True, data_only=False)
    try:
        if 'Reported Data' not in values.sheetnames:
            raise ValueError('Workbook has no Reported Data sheet.')
        rows = list(values['Reported Data'].values)
        raw_rows = list(formulas['Reported Data'].values)
        headers = [str(v).strip() if v is not None else '' for v in rows[0]]
        required = [CARRIER, MOBILITY]
        if any(name not in headers for name in required):
            raise ValueError(f'Required headers: {required}')
        resistivity = [i for i, h in enumerate(headers) if h.startswith('Resistivity (')]
        if len(resistivity) != 1:
            raise ValueError('Expected one resistivity column with units in its header.')
        rho_col = resistivity[0]
        unit = headers[rho_col].lower()
        if 'mohm' in unit or 'mω' in unit:
            factor = 1
        elif 'ohm' in unit or 'ω' in unit:
            factor = 1000
        else:
            raise ValueError(f'Unrecognised resistivity unit: {headers[rho_col]}')
        n_col, mu_col = headers.index(CARRIER), headers.index(MOBILITY)
        records = []
        skipped = 0
        for row_number, row in enumerate(rows[1:], start=2):
            if not any(v is not None for v in row):
                continue
            n, mu = positive_number(row[n_col]), positive_number(row[mu_col])
            if n is None or mu is None:
                skipped += 1
                continue
            rho = positive_number(row[rho_col])
            raw = raw_rows[row_number - 1][rho_col]
            origin = 'workbook formula (cached)' if isinstance(raw, str) and raw.startswith('=') else 'workbook value'
            if rho is not None:
                rho *= factor
            elif derive_missing and raw is None:
                rho = isoresistivity_mobility(n, mu)
                origin = 'derived by script (single-carrier relation)'
            else:
                origin = 'missing or invalid resistivity'
            def source(name):
                return row[headers.index(name)] if name in headers else ''
            records.append({'observation': source('Observation ID'), 'paper': source('Paper ID'),
                            'source_row': row_number, 'n_cm-3': n, 'mobility_cm2_Vs': mu,
                            'resistivity_mohm_cm': rho, 'resistivity_origin': origin})
        if not records:
            raise ValueError('No observations have positive, finite carrier concentration and mobility.')
        return records, skipped
    finally:
        values.close()
        formulas.close()


def generate(workbook, output_dir, levels, derive_missing=False, label_points=False):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.colors import LogNorm
    from matplotlib.cm import ScalarMappable

    records, skipped = read_observations(workbook, derive_missing)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({'font.size': 13, 'axes.labelsize': 14, 'svg.fonttype': 'none'})
    fig, ax = plt.subplots(figsize=(9.5, 6.5), layout='constrained')
    xs = [r['n_cm-3'] for r in records]
    ys = [r['mobility_cm2_Vs'] for r in records]
    xmin, xmax = min(xs) / 1.8, max(xs) * 1.8
    ymin, ymax = min(ys) / 1.6, max(ys) * 1.6
    ax.set(xscale='log', yscale='log', xlim=(xmin, xmax), ylim=(ymin, ymax),
           xlabel=r'Carrier concentration, $n$ (cm$^{-3}$)',
           ylabel=r'Mobility, $\mu$ (cm$^2$ V$^{-1}$ s$^{-1}$)',
           title='Figure 1. Mobility versus carrier concentration')
    coloured = [r for r in records if r['resistivity_mohm_cm'] is not None]
    if coloured:
        rho_values = [r['resistivity_mohm_cm'] for r in coloured]
        low = 10 ** math.floor(math.log10(min(rho_values)))
        high = 10 ** math.ceil(math.log10(max(rho_values)))
        if low == high:
            low /= math.sqrt(10)
            high *= math.sqrt(10)
        norm = LogNorm(low, high)
        groups = [('workbook value', 'o', 'Workbook value'),
                  ('workbook formula (cached)', '^', 'Calculated in workbook'),
                  ('derived by script (single-carrier relation)', 's', 'Calculated by script')]
        for origin, marker, label in groups:
            group = [r for r in coloured if r['resistivity_origin'] == origin]
            if group:
                ax.scatter([r['n_cm-3'] for r in group], [r['mobility_cm2_Vs'] for r in group],
                           c=[r['resistivity_mohm_cm'] for r in group], cmap='viridis', norm=norm,
                           marker=marker, s=85, edgecolors='black', linewidths=.7, label=label, zorder=3)
        cbar = fig.colorbar(ScalarMappable(norm=norm, cmap='viridis'), ax=ax, pad=.025)
        cbar.set_label(r'Resistivity, $\rho$ (m$\Omega$ cm), logarithmic scale')
    missing = [r for r in records if r['resistivity_mohm_cm'] is None]
    if missing:
        ax.scatter([r['n_cm-3'] for r in missing], [r['mobility_cm2_Vs'] for r in missing],
                   color='grey', marker='x', s=70, label='Resistivity unavailable', zorder=3)

    # Constant resistivity has slope -1 in log(n)-log(mu) coordinates.
    for rho in levels:
        left = max(xmin, 1000 / (Q * rho * ymax))
        right = min(xmax, 1000 / (Q * rho * ymin))
        if left >= right:
            continue
        ax.plot([left, right], [isoresistivity_mobility(left, rho), isoresistivity_mobility(right, rho)],
                linestyle=':', color='0.45', linewidth=1.25, zorder=1)
        label_x = left * (right / left) ** .82
        ax.text(label_x, isoresistivity_mobility(label_x, rho), f'{rho:g} mΩ cm',
                fontsize=10, color='0.3', ha='center', va='center', zorder=2,
                bbox={'facecolor': 'white', 'edgecolor': 'none', 'alpha': .85, 'pad': 1})
    if label_points:
        for r in records:
            ax.annotate(str(r['observation']), (r['n_cm-3'], r['mobility_cm2_Vs']),
                        xytext=(5, 6), textcoords='offset points', fontsize=9)
    ax.tick_params(which='both', direction='out', top=True, right=True)
    ax.legend(loc='upper left', fontsize=10, framealpha=.95)
    for extension in ['png', 'svg']:
        fig.savefig(output_dir / f'figure1.{extension}', dpi=300)
    plt.close(fig)
    with (output_dir / 'figure1-data.csv').open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    (output_dir / 'figure1-notes.md').write_text(
        f'# Figure 1 provenance\n\nSource: {Path(workbook).name}, Reported Data.\n\n'
        f'Plotted: {len(records)} observations; omitted for missing/nonpositive n or mobility: {skipped}.\n'
        f'Colour available for {len(coloured)} observations. Resistivity origin is retained in the CSV.\n\n'
        'Dotted contours use μ = 1000/(q n ρ), with ρ in mΩ cm, n in cm⁻³, '
        'and μ in cm² V⁻¹ s⁻¹; q = 1.602176634 × 10⁻¹⁹ C. '
        'This is a single-carrier conductivity relation, not an independent measurement. '
        'Workbook formula values are read from saved Excel caches; recalculate and save the workbook if needed. '
        'The script never edits the workbook.\n', encoding='utf-8')
    print(f'Plotted {len(records)} observations ({len(coloured)} with colour); skipped {skipped}.')
    print(f'Saved {output_dir.resolve() / "figure1.png"} and SVG, data CSV, provenance notes.')
    return records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'example-output')
    parser.add_argument('--iso', type=float, nargs='+', default=[.1, .2, .5, 1, 2, 5, 10],
                        help='constant resistivity levels in mΩ cm')
    parser.add_argument('--derive-missing', action='store_true', help='calculate only genuinely blank resistivity cells from n and mobility')
    parser.add_argument('--label-points', action='store_true', help='label points with Observation IDs')
    args = parser.parse_args()
    if any(not math.isfinite(v) or v <= 0 for v in args.iso):
        parser.error('--iso values must be positive and finite')
    try:
        generate(args.workbook, args.output_dir, args.iso, args.derive_missing, args.label_points)
    except (ValueError, OSError) as exc:
        parser.exit(1, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
