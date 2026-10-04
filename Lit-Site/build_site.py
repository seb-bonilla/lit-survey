"""Export a literature workbook's Main sheet to a MkDocs homepage."""
from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
FIGURE = re.compile(r"Figure\s+(\d+)", re.IGNORECASE)
ASSET_TYPES = {".png", ".jpg", ".jpeg", ".svg", ".webp", ".gif", ".html"}


def text(value):
    return "" if value is None else str(value).strip()


def read_main(workbook: Path, sheet_name: str):
    """Read source cells without editing the workbook or dropping formulas silently."""
    wb = load_workbook(workbook, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(max_col=3):
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"{sheet_name}!{cell.coordinate} is a formula. Use narrative text in columns A-C; export does not calculate formulas.")
            rows.append(tuple(text(cell.value) for cell in row))
        if not rows or not rows[0][0]:
            raise ValueError(f"{sheet_name}!A1 must contain the page title.")
        if rows[0][1] or rows[0][2]:
            raise ValueError(f"Keep {sheet_name}!B1 and C1 empty; A1 is reserved for the title.")
        return rows[0][0], rows[1:]
    finally:
        wb.close()


def figure_numbers(marker: str, row: int):
    if not marker:
        return []
    result = []
    for part in re.split(r"[,;\n]+", marker):
        match = FIGURE.fullmatch(part.strip())
        if not match:
            raise ValueError(f"Main row {row}: invalid figure marker {part!r}; use Figure 1, Figure 2, etc.")
        result.append(int(match.group(1)))
    return result


def load_figures(mapping: Path | None, wanted: set[int]):
    if not wanted:
        return {}
    if mapping is None:
        raise ValueError("Column C contains figure markers. Supply --figures path/to/figures.json, or use --skip-figures for a narrative-only export.")
    data = json.loads(mapping.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict):
        raise ValueError("Figure mapping must be a JSON object keyed by Figure 1, Figure 2, etc.")
    figures = {}
    for number in sorted(wanted):
        key = f"Figure {number}"
        entry = data.get(key)
        if not isinstance(entry, dict) or not isinstance(entry.get("file"), str):
            raise ValueError(f"Missing file entry for {key} in {mapping}")
        source = Path(entry["file"]).expanduser()
        if not source.is_absolute():
            source = mapping.parent / source
        source = source.resolve()
        if not source.is_file() or source.suffix.lower() not in ASSET_TYPES:
            raise ValueError(f"Missing or unsupported asset for {key}: {source}")
        figures[number] = {"source": source, "caption": text(entry.get("caption")) or key,
                           "relative": f"assets/figures/figure{number}{source.suffix.lower()}"}
    return figures


def render_main(title, rows, figures, skip_figures=False):
    # Sequential rendering retains row order, including section intros and C-only rows.
    lines = [f"# {title}", ""]
    for row_number, (heading, paragraph, marker) in enumerate(rows, start=2):
        if heading:
            lines.extend([f"{'###' if paragraph else '##'} {heading}", ""])
        if paragraph:
            lines.extend([paragraph, ""])
        for number in figure_numbers(marker, row_number):
            if skip_figures:
                lines.extend([f"*Figure {number} omitted from this narrative-only export.*", ""])
                continue
            figure = figures[number]
            caption = figure["caption"]
            url = quote(figure["relative"], safe="/")
            if figure["source"].suffix.lower() == ".html":
                lines.extend([f'<iframe src="{url}" title="{html.escape(caption, quote=True)}" style="width:100%;height:650px;border:0" loading="lazy"></iframe>', ""])
            else:
                alt = caption.replace("[", r"\[").replace("]", r"\]").replace("\n", " ")
                lines.extend([f"![{alt}]({url})", ""])
            lines.extend([caption, ""])
    return "\n".join(lines).rstrip() + "\n"


def export(workbook, sheet="Main", docs_dir=None, mapping=None, skip_figures=False):
    workbook = Path(workbook).expanduser().resolve()
    if not workbook.is_file():
        raise ValueError(f"Workbook not found: {workbook}")
    title, rows = read_main(workbook, sheet)
    wanted = {n for row_number, row in enumerate(rows, start=2) for n in figure_numbers(row[2], row_number)}
    figures = {} if skip_figures else load_figures(Path(mapping).resolve() if mapping else None, wanted)
    markdown = render_main(title, rows, figures, skip_figures)
    docs = Path(docs_dir).resolve() if docs_dir else ROOT / "docs"
    # All input validation completes before modifying generated outputs.
    docs.mkdir(parents=True, exist_ok=True)
    for figure in figures.values():
        destination = docs / figure["relative"]
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.resolve() != figure["source"]:
            shutil.copy2(figure["source"], destination)
    target = docs / "index.md"
    temporary = target.with_suffix(".md.tmp")
    temporary.write_text(markdown, encoding="utf-8")
    temporary.replace(target)
    print(f"Generated {target} from {sheet!r}: {len(rows)} rows, {len(figures)} figure assets.")
    if skip_figures and wanted:
        print(f"WARNING: omitted {len(wanted)} referenced figures; review placeholders before publishing.")
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="path to the private XLSX workbook")
    parser.add_argument("--sheet", default="Main")
    parser.add_argument("--figures", type=Path, help="JSON mapping from Figure N to file and caption")
    parser.add_argument("--skip-figures", action="store_true", help="explicitly export narrative with omission notes")
    parser.add_argument("--docs-dir", type=Path, help="optional export destination; default Lit-Site/docs")
    parser.add_argument("--build", action="store_true", help="also build Lit-Site's static HTML with MkDocs")
    args = parser.parse_args()
    if args.build and args.docs_dir:
        parser.error("--build uses the standard docs folder; use MkDocs separately with custom --docs-dir")
    try:
        export(args.workbook, args.sheet, args.docs_dir, args.figures, args.skip_figures)
        if args.build:
            subprocess.run([sys.executable, "-m", "mkdocs", "build", "--strict", "-f", str(ROOT / "mkdocs.yml")], check=True)
    except (ValueError, OSError, KeyError, subprocess.CalledProcessError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
